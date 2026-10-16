"""Read official workbooks without editing them. Preserve cell-level evidence.
Paired arrival/departure rows distinguish passing times from passenger stops.
Unrecognized layouts and timetable-internal station names are NOT auto-approved.
"""
import datetime, hashlib, json, pathlib, re, shutil, sys
import openpyxl
from openpyxl.utils import get_column_letter
sys.stdout.reconfigure(encoding='utf-8')
directory, board_file, output = map(pathlib.Path, sys.argv[1:4])
output.mkdir(parents=True, exist_ok=True)
board = {str(r['bdIdx']): r for r in json.loads(board_file.read_text(encoding='utf-8'))['boardList']}
trips, documents, skipped = [], [], []
clean = lambda v: re.sub(r'\s+', '', str(v or ''))


def time_value(v):
    if isinstance(v, datetime.time):
        return v.isoformat()
    if isinstance(v, datetime.timedelta):
        return str(v)
    if isinstance(v, str) and re.fullmatch(r'\d{1,2}:\d{2}(?::\d{2})?', v.strip()):
        return v.strip()
    return None


for file in sorted(directory.glob('*.xlsx')):
    row = board[file.stem]
    book = openpyxl.load_workbook(file, data_only=True, read_only=True)
    destination = output / file.name
    if file.resolve() != destination.resolve():
        shutil.copyfile(file, destination)
    match = re.search(r'(202\d)\s*\.\s*(\d{1,2})\s*\.\s*(\d{1,2})', row['bdTitle'])
    effective = datetime.date(*map(int, match.groups())).isoformat() if match else None
    source_id = 'korail-timetable-' + file.stem
    documents.append({'id': source_id, 'kind': 'official-timetable', 'publisher': '한국철도공사', 'title': row['bdTitle'], 'url': 'https://www.korail.com/file/cubedata/COMMON/'+row['fileId'][0], 'validFrom': effective, 'retrievedOn': datetime.date.today().isoformat(), 'license': 'unverified', 'licenseStatus': 'pending', 'files': [{'path': 'source/timetables/'+file.name, 'sha256': hashlib.file_digest(file.open('rb'), 'sha256').hexdigest()}]})
    for sheet in book:
        values = list(sheet.iter_rows())
        header = next((i for i, r in enumerate(values[:10]) if clean(r[0].value) == '열차번호'), None)
        origin_row = next((i for i, r in enumerate(values[:10]) if clean(r[0].value) == '시발역'), None)
        dest_row = next((i for i, r in enumerate(values[:10]) if clean(r[0].value) == '종착역'), None)
        if header is None or origin_row is None or dest_row is None:
            skipped.append({'sourceId': source_id, 'sheet': sheet.title, 'reason': 'UNSUPPORTED_LAYOUT_REQUIRES_EXPLICIT_MAPPING'})
            continue
        station_rows = []
        for i in range(header+1, len(values)-1):
            name = clean(values[i][0].value)
            if re.fullmatch(r'[가-힣A-Za-z0-9()·]+', name) and not clean(values[i+1][0].value) and any(time_value(c.value) for c in values[i][1:]+values[i+1][1:]):
                station_rows.append((i, name))
        if len(station_rows) < 2:
            skipped.append({'sourceId': source_id, 'sheet': sheet.title, 'reason': 'NO_PAIRED_STATION_ROWS'})
            continue
        for col in range(1, sheet.max_column):
            number = clean(values[header][col].value)
            if not re.fullmatch(r'[A-Za-z]*\d+', number):
                continue
            origin, terminal = clean(values[origin_row][col].value), clean(values[dest_row][col].value)
            calls, passes = [], []
            for i, name in station_rows:
                arrival, departure = time_value(values[i][col].value), time_value(values[i+1][col].value)
                if not arrival and not departure:
                    continue
                record = {'operationalStationName': name, 'arrival': arrival, 'departure': departure, 'sourceCells': [get_column_letter(col+1)+str(i+1), get_column_letter(col+1)+str(i+2)]}
                # A departure-only time at an intermediate station is a passing
                # time in these workbooks; do not turn it into a typing stop.
                if arrival or name == origin:
                    calls.append(record)
                else:
                    passes.append(record)
            if len(calls)<2:
                continue
            trips.append({'id': source_id+':'+sheet.title+':'+number, 'sourceId': source_id, 'sheet': sheet.title, 'trainNumber': number, 'trainNumberCell': values[header][col].coordinate, 'calendarLabel': sheet.title+' / '+row['bdTitle'], 'validFrom': effective, 'originOperationalName': origin, 'destinationOperationalName': terminal, 'calls': calls, 'passing': passes, 'status': 'candidate', 'reviewRequired': ['OPERATIONAL_NAMES_TO_OFFICIAL_STATIONS', 'CALENDAR_AND_EXCEPTIONS', 'SERVICE_TRACK_AND_DIRECTION', 'SOURCE_LICENSE']})
    book.close()
result = {'status': 'INCOMPLETE', 'documents': documents, 'tripCandidates': trips, 'unparsedSheets': skipped}
(output.parent.parent/'timetable-candidates.json').write_text(json.dumps(result,ensure_ascii=False,separators=(',',':'))+'\n',encoding='utf-8')
registry = output.parent.parent/'sources.json'
existing = json.loads(registry.read_text(encoding='utf-8'))
ids = {d['id'] for d in documents}
registry.write_text(json.dumps([s for s in existing if s['id'] not in ids]+documents,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'documents':len(documents),'tripCandidates':len(trips),'unparsedSheets':len(skipped)},ensure_ascii=False))
